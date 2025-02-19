import requests


def get_weather_data(city):
    url = f"http://api.openweathermap.org/data/2.5/weather?q={city}&appid=YOUR_API_KEY"
    response = requests.get(url)
    return response.json()

def get_country_info(country_code):
    url = f"https://restcountries.com/v3.1/all"
    response = requests.get(url)
    countries = response.json()
    for country in countries:
        if country["cca2"] == country_code:
            return country
    return None

def combine_data(weather_data, country_data):
    combined = {
        'City': weather_data['name'],
        'Country': country_data['name'],
        'Temperature': weather_data['main']['temp'],
        'Weather': weather_data['weather'][0]['description']
    }
    return combined

from openpyxl import Workbook

def write_to_excel(data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Combined Data"
    
    # Add headers
    headers = ['City', 'Country', 'Temperature', 'Weather']
    ws.append(headers)
    
    # Add the data
    for entry in data:
        ws.append([entry['City'], entry['Country'], entry['Temperature'], entry['Weather']])
    
    # Save the file
    wb.save("combined_data.xlsx")

from openpyxl.chart import BarChart, Reference

def add_chart_to_excel(wb):
    ws = wb.active
    chart = BarChart()
    data = Reference(ws, min_col=3, min_row=2, max_row=5, max_col=3)
    categories = Reference(ws, min_col=1, min_row=2, max_row=5)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    ws.add_chart(chart, "E5")
    write_to_excel(combine_data)

